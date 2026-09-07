# -*- coding: utf-8 -*-
"""Génère docs/PathFinder-essentiel.pptx.

Diaporama court : les essentiels du « PathFinder » de dotnet-new-scan
(PathFinder.ScanMvc) — recherche de plus court chemin sur dbo.LINE_VIS_EDG :
modèle de données, graphe chargé en mémoire (CSR), les quatre algorithmes
(BFS bidirectionnel / Dijkstra / Dijkstra bidirectionnel / A* ALT), le
pipeline de décision, les durées mesurées, et les fichiers clés.

Même charte visuelle que docs/generer_ppt_recherche_chemin.py (navy + vert).
python-pptx requis (pip install python-pptx).

Lancer :  python docs/generer_ppt_pathfinder_essentiel.py
Sortie  :  docs/PathFinder-essentiel.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DOCS = Path(r"C:\Users\amami\GitHub\restitutiondonnees") / "docs"
OUT = DOCS / "PathFinder-essentiel.pptx"

# --- palette -------------------------------------------------------------
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
AMBER = RGBColor(0xC9, 0x78, 0x2D)
INK = RGBColor(0x22, 0x2A, 0x35)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF1, 0xF3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


# --- aides (reprises de generer_ppt_recherche_chemin.py) ----------------
def _box(slide, x, y, w, h):
    return slide.shapes.add_textbox(x, y, w, h)


def slide_header(slide, title, kicker=None):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(0.16))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tb = _box(slide, Inches(0.6), Inches(0.34), SW - Inches(1.2), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    if kicker:
        r = tf.paragraphs[0].add_run()
        r.text = kicker.upper()
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = GREEN
        p = tf.add_paragraph()
    else:
        p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY


def new_slide(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        slide_header(s, title, kicker)
    return s


def bullets(slide, items, x=Inches(0.7), y=Inches(1.75), w=None, h=None, size=16, gap=8):
    w = w or (SW - Inches(1.4))
    h = h or (SH - y - Inches(0.4))
    tf = _box(slide, x, y, w, h).text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        lvl, text = it if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = ("•  " if lvl == 0 else "–  ") + text
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.color.rgb = INK if lvl == 0 else GREY


def table(slide, rows, x=Inches(0.7), y=Inches(1.8), w=None, h=None, font=12, col_widths=None):
    w = w or (SW - Inches(1.4))
    h = h or Inches(0.42 * len(rows))
    tbl = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
    if col_widths:
        tot = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(w * cw / tot))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.margin_left = c.margin_right = Pt(7)
            c.margin_top = c.margin_bottom = Pt(3)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.text = str(val)
            for run in c.text_frame.paragraphs[0].runs:
                run.font.size = Pt(font)
                if ri == 0:
                    run.font.bold = True
                    run.font.color.rgb = WHITE
                else:
                    run.font.color.rgb = INK
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if ri == 0 else (WHITE if ri % 2 else LIGHT)
    return tbl


def code(slide, text, x=Inches(0.7), y=Inches(1.8), w=None, h=None, size=12):
    w = w or (SW - Inches(1.4))
    h = h or (SH - y - Inches(0.4))
    box = slide.shapes.add_shape(1, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF6, 0xF8, 0xFA)
    box.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
    box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(12)
    tf.margin_top = tf.margin_bottom = Pt(10)
    for i, line in enumerate(text.strip("\n").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line or " "
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        r.font.color.rgb = INK


def footnote(slide, text):
    tb = _box(slide, Inches(0.7), SH - Inches(0.55), SW - Inches(1.4), Inches(0.45))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = GREY


# =======================================================================
# 1. TITRE
# =======================================================================
s = new_slide()
tb = _box(s, Inches(0.9), Inches(2.5), SW - Inches(1.8), Inches(2.6))
tf = tb.text_frame
tf.word_wrap = True
r = tf.paragraphs[0].add_run()
r.text = "PathFinder — l'essentiel"
r.font.size = Pt(44)
r.font.bold = True
r.font.color.rgb = NAVY
p = tf.add_paragraph()
r = p.add_run()
r.text = ("dotnet-new-scan / PathFinder.ScanMvc\n"
          "Recherche de plus court chemin sur dbo.LINE_VIS_EDG :\n"
          "graphe en mémoire + quatre algorithmes de parcours")
r.font.size = Pt(18)
r.font.color.rgb = GREY
footnote(s, "ASP.NET Core MVC (Razor) · C# 8.0 / .NET 7.0 · aucun JavaScript")

# =======================================================================
# 2. LE PROBLÈME
# =======================================================================
s = new_slide("Le problème", "Ce que fait PathFinder")
bullets(s, [
    "Question : existe-t-il un chemin orienté entre deux nœuds du graphe, et "
    "quel est le plus court ?",
    "Le graphe : la table SQL Server dbo.LINE_VIS_EDG (jeu de démo : "
    "100 008 nœuds, 400 787 arêtes).",
    "Un nœud = une valeur qui apparaît en colonne Nodes ou NodesLie. Pas de "
    "table de nœuds séparée.",
    "Résultat affiché : la chaîne source → … → cible, et la transformation "
    "portée par chaque arête.",
    (1, "Formulaire GET, URL partageable, résultat mis en cache 5 minutes."),
    (1, "Rendu 100 % côté serveur (Razor + SVG) — pas de JavaScript, pas d'API JSON."),
])

# =======================================================================
# 3. MODÈLE DE DONNÉES
# =======================================================================
s = new_slide("Modèle de données", "dbo.LINE_VIS_EDG")
table(s, [
    ("Colonne", "Type", "Rôle"),
    ("Nodes", "VARCHAR(8000)", "un des deux nœuds"),
    ("NodesLie", "VARCHAR(8000)", "l'autre nœud"),
    ("Direction", "VARCHAR", "'predecesseur' ou 'successeur'"),
    ("Transformation", "VARCHAR", "libellé porté par l'arête"),
], y=Inches(1.8), h=Inches(2.4), col_widths=[2, 2, 5])
code(s,
     "Direction = 'predecesseur'  ->  arête  Nodes  -> NodesLie\n"
     "Direction = 'successeur'    ->  arête  NodesLie -> Nodes\n\n"
     "// paramètres SQL typés VARCHAR explicitement (sinon NVARCHAR =>\n"
     "// conversion de colonne => scan au lieu de seek d'index)",
     y=Inches(4.5), h=Inches(2.2), size=13)

# =======================================================================
# 4. LE GRAPHE EN MÉMOIRE (CSR)
# =======================================================================
s = new_slide("Le graphe en mémoire", "§ 11.7 — structure CSR")
bullets(s, [
    "Au démarrage, tout le graphe est chargé en RAM en tableaux d'entiers "
    "(format CSR, Compressed Sparse Row).",
    "Chargement sur un Thread d'arrière-plan (pas d'async) : le serveur répond "
    "tout de suite ; repli sur le BFS SQL tant que le graphe n'est pas prêt.",
    "~12 Mo pour 100 000 nœuds ; chargement ~1,2–2,0 s.",
], y=Inches(1.7), h=Inches(2.3), size=15)
code(s,
     "_names[i]     : indice -> nom du nœud            (n entrées)\n"
     "_index[nom]   : nom -> indice\n"
     "_fwdOffset[u].._fwdOffset[u+1]  -> tranche de _fwdTarget = successeurs de u\n"
     "_revOffset[u].._revOffset[u+1]  -> tranche de _revSource = prédécesseurs de u\n\n"
     "Successors(u) = _fwdTarget.AsSpan(...)   // aucune allocation, aucune copie",
     y=Inches(4.3), h=Inches(2.5), size=12)

# =======================================================================
# 5. LES QUATRE ALGORITHMES
# =======================================================================
s = new_slide("Quatre algorithmes", "Sélection par le menu « Algorithme »")
table(s, [
    ("algo=", "Nom", "Principe", "Correct si poids ≠ 1"),
    ("bfs", "BFS bidirectionnel  (défaut)", "deux fronts en largeur qui se rejoignent au milieu", "non"),
    ("dijkstra", "Dijkstra", "file de priorité, un seul front", "oui"),
    ("dijkstra-bi", "Dijkstra bidirectionnel", "deux Dijkstra en miroir, arrêt sur topF+topB ≥ best", "oui"),
    ("astar", "A* (heuristique ALT)", "Dijkstra guidé par des repères précalculés", "oui"),
], y=Inches(2.0), col_widths=[1.3, 2.6, 5.5, 1.6], font=12)
bullets(s, [
    "Poids d'arêtes tous égaux à 1 : les quatre renvoient le même chemin "
    "(à un ex æquo près). Seule la durée change.",
], y=Inches(5.2), size=14)

# =======================================================================
# 6. BFS BIDIRECTIONNEL
# =======================================================================
s = new_slide("BFS bidirectionnel", "Le choix par défaut")
code(s,
     "        source ──▶─▶─▶ ┐        ┌ ◀─◀─◀── cible\n"
     "  front avant (successeurs)   front arrière (prédécesseurs)\n"
     "               ╲            ╱\n"
     "                ╲  point   ╱\n"
     "                 ╲  de    ╱\n"
     "                  ╲ jonction\n"
     "                   ▼\n"
     "        recoller les deux demi-chemins (BuildPath)",
     y=Inches(1.8), h=Inches(3.0), size=13)
bullets(s, [
    "≈ 2·b^(R/2) nœuds visités au lieu de b^R (b = degré ≈ 4, R = longueur).",
    "File FIFO, aucun tas, aucun précalcul. Garde-fou : 300 000 nœuds par sens.",
    "Devient FAUX si les arêtes deviennent pondérées.",
], y=Inches(5.0), size=14)

# =======================================================================
# 7. A* / ALT
# =======================================================================
s = new_slide("A* et l'heuristique ALT", "A* sans coordonnées")
bullets(s, [
    "A* classe les nœuds par g(n) + h(n) : g = coût déjà parcouru, h = "
    "estimation du coût restant vers la cible.",
    "h doit être admissible (ne jamais surestimer) pour rester optimal.",
    "Pas de coordonnées ici → heuristique ALT (A*, Landmarks, Triangle "
    "inequality) :",
    (1, "PrepareLandmarks choisit 6 repères bien étalés (farthest-first)."),
    (1, "Pour chaque repère : distance BFS vers / depuis tous les nœuds, précalculée au démarrage (~5 Mo)."),
    (1, "L'inégalité triangulaire donne une borne inférieure valable de dist(n, cible)."),
    "Sur ce graphe aléatoire, l'heuristique est faible : A* reste lent. Utile "
    "surtout sur un graphe à structure (couches, coordonnées).",
], y=Inches(1.7), size=14)

# =======================================================================
# 8. PIPELINE DE DÉCISION
# =======================================================================
s = new_slide("Avant de lancer un parcours", "HomeController — ordre des vérifications")
code(s,
     "1.  Condensation SCC (§ 11.5, si calculée)\n"
     "        NotReachable  ->  « aucun chemin »   (verdict EXACT, sans parcours)\n"
     "\n"
     "2.  sinon : Composantes faibles (§ 11.4, si calculées)\n"
     "        îles différentes  ->  « aucun chemin »   (sans parcours)\n"
     "\n"
     "3.  sinon : Cache applicatif (5 min)  ->  résultat déjà connu ?\n"
     "\n"
     "4.  sinon : Parcours avec l'algo choisi, sur le graphe en mémoire\n"
     "\n"
     "5.  graphe pas encore chargé  ->  repli BFS SQL palier par palier",
     y=Inches(1.8), h=Inches(4.4), size=14)
footnote(s, "Les deux pré-calculs (§ 11.4 / § 11.5) se déclenchent depuis la page /Scan.")

# =======================================================================
# 9. DURÉES MESURÉES
# =======================================================================
s = new_slide("Durées mesurées", "ms par appel — graphe en mémoire, hors cache")
table(s, [
    ("Algorithme", "Chemin existant (médiane)", "Aucun chemin"),
    ("BFS bidirectionnel", "0,06 ms", "22,8 ms"),
    ("Dijkstra bidirectionnel", "0,37 ms", "0,001 ms"),
    ("A* (ALT)", "8,3 ms", "103,7 ms"),
    ("Dijkstra", "22,4 ms", "61,8 ms"),
], y=Inches(2.0), col_widths=[3, 3, 3], font=14)
bullets(s, [
    "Le bidirectionnel écrase le reste (deux fronts de rayon R/2).",
    "Dijkstra bidirectionnel : imbattable sur « aucun chemin » — un front se "
    "vide instantanément.",
    "A* : le pire sur « aucun chemin » (calcule l'heuristique sans jamais élaguer).",
], y=Inches(4.6), size=14)

# =======================================================================
# 10. VERDICT
# =======================================================================
s = new_slide("Verdict pour ce projet", "Quel algorithme ?")
table(s, [
    ("Situation", "Choix"),
    ("Poids = 1 (actuel)", "BFS bidirectionnel — rien ne le bat, aucun précalcul"),
    ("Si Transformation devient un coût", "Dijkstra bidirectionnel — quasi la même vitesse, correct pondéré"),
    ("Graphe avec structure (couches, coord.)", "A* — sinon son précalcul n'est pas rentable"),
    ("« Aucun chemin » à trancher", "Condensation SCC — O(1), avant tout parcours"),
], y=Inches(2.0), col_widths=[3.5, 6], font=13)

# =======================================================================
# 11. CONTRAINTES
# =======================================================================
s = new_slide("Contraintes techniques", "Cible imposée")
bullets(s, [
    "C# 8.0 / .NET 7.0, sans async / await (préchargement sur un Thread).",
    "ImplicitUsings désactivé ; pas de record, pas de namespaces à portée de "
    "fichier, pas de collection expressions [].",
    "Microsoft.Data.SqlClient 5.2.2 (les 7.x ne ciblent que .NET 8+).",
    "Requêtes SQL toutes paramétrées ; POST du /Scan protégés par jeton "
    "anti-forgery.",
    "Aucun SDK .NET 7 requis : le SDK 10 compile pour net7.0, le runtime 7.0 "
    "exécute.",
], y=Inches(1.8), size=15)

# =======================================================================
# 12. FICHIERS CLÉS
# =======================================================================
s = new_slide("Fichiers clés", "Le cœur PathFinder")
table(s, [
    ("Groupe", "Fichiers"),
    ("Mise en mémoire des données",
     "LineVisEdgRepository.cs (StreamAllDirectedEdges) · DirectedGraph.cs (Build → CSR) · "
     "InMemoryGraphService.cs (+ GraphPreloader)"),
    ("PathFinder (.cs)",
     "DirectedGraph.cs (BFS bi, Dijkstra, Dijkstra bi, A*/ALT) · InMemoryGraphService.cs (passe-plats) · "
     "HomeController.cs (aiguillage + cache) · PathViewModel.cs · LineVisEdgRepository.cs (BFS SQL de repli, DescribePath)"),
    ("Front minimal (HTML rendu serveur)",
     "Views/Home/Index.cshtml (formulaire + résultat) · Views/Shared/_Layout.cshtml · wwwroot/css/site.css"),
], y=Inches(1.8), col_widths=[2.4, 8], font=11)
footnote(s, "Détail : dotnet-new-scan/README.md · docs/algorithmes-chemin.md · "
            "docs/Specification-fonctionnelle-dotnet-new-scan.docx")

# =======================================================================
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"OK -> {OUT}  ({sum(1 for _ in prs.slides)} diapos)")
